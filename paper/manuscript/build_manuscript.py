#!/usr/bin/env python3
"""Build the submission DOCX (and, via LibreOffice, PDF) from manuscript.md.

Source dialect (deliberately small):
  # Title                         first line
  AUTHOR: / AFFIL: / EMAIL: / ORCID:   front-matter lines
  ## Section / ### Subsection
  **Keywords:** ...               keywords line
  TABLE1                          placeholder where Table 1 is inserted
  ## Table  ...  TABLE1: caption  followed by a pipe table
  ## Figure captions              ![Fig. N. caption](figures/x.png) blocks
  Inline: *italic*, **bold**, `code`, ^{sup}, _{sub}; Unicode elsewhere.
  Bracketed [PROPOSED, author to confirm: ...] passages are kept verbatim
  and rendered in a highlighted run so they cannot be missed.

Usage: build_manuscript.py [--figures inline|end] [--stem NAME] [--pdf]
  --figures end    (default) figures + captions collected after References
                   (journal DOCX layout)
  --figures inline figures placed after the paragraph that first mentions
                   "Fig. N" (preprint layout)
  --pdf            also convert to PDF with LibreOffice (soffice --headless)
"""
import argparse
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for clarity)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "manuscript.md")


# ----------------------------------------------------------------- parsing
def parse(md):
    lines = md.splitlines()
    doc = {"title": None, "meta": {}, "blocks": [], "table": None, "figures": []}
    i = 0
    cur_section = None
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# ") and doc["title"] is None:
            doc["title"] = ln[2:].strip(); i += 1; continue
        m = re.match(r"^(AUTHOR|AFFIL|EMAIL|ORCID):\s*(.*)$", ln)
        if m:
            doc["meta"][m.group(1)] = m.group(2).strip(); i += 1; continue
        if ln.startswith("## "):
            cur_section = ln[3:].strip()
            if cur_section not in ("Table", "Figure captions"):
                doc["blocks"].append(("h1", cur_section))
            i += 1; continue
        if ln.startswith("### "):
            doc["blocks"].append(("h2", ln[4:].strip())); i += 1; continue
        if cur_section == "Table":
            m = re.match(r"^TABLE1:\s*(.*)$", ln)
            if m:
                caption = m.group(1).strip()
                rows = []
                i += 1
                while i < len(lines) and lines[i].startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                        rows.append(cells)
                    i += 1
                doc["table"] = {"caption": caption, "rows": rows}
                continue
            i += 1; continue
        if cur_section == "Figure captions":
            m = re.match(r"^!\[(.*)\]\((.*)\)\s*$", ln)
            if m:
                doc["figures"].append({"caption": m.group(1), "path": m.group(2)})
            i += 1; continue
        if ln.strip() == "TABLE1":
            doc["blocks"].append(("table", None)); i += 1; continue
        if ln.strip() == "":
            i += 1; continue
        # paragraph: gather until blank line
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() != "" and not lines[i].startswith("#"):
            para.append(lines[i]); i += 1
        text = " ".join(p.strip() for p in para)
        if text.startswith("**Keywords:**"):
            doc["blocks"].append(("keywords", text[len("**Keywords:**"):].strip()))
        else:
            doc["blocks"].append(("p", text))
    return doc


# --------------------------------------------------------------- inline runs
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^\{[^}]*\}|_\{[^}]*\}|\[PROPOSED[^\]]*\])")


def add_inline(par, text, base_italic=False, base_bold=False, size=None):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], base_italic, base_bold, size=size)
        tok = m.group(0)
        if tok.startswith("**"):
            _run(par, tok[2:-2], base_italic, True, size=size)
        elif tok.startswith("*"):
            _run(par, tok[1:-1], True, base_bold, size=size)
        elif tok.startswith("`"):
            _run(par, tok[1:-1], base_italic, base_bold, font="Consolas", size=size)
        elif tok.startswith("^{"):
            _run(par, tok[2:-1], base_italic, base_bold, sup=True, size=size)
        elif tok.startswith("_{"):
            _run(par, tok[2:-1], base_italic, base_bold, sub=True, size=size)
        elif tok.startswith("[PROPOSED"):
            _run(par, tok, base_italic, base_bold, highlight=True, size=size)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], base_italic, base_bold, size=size)


def _run(par, s, italic=False, bold=False, font=None, sup=False, sub=False, highlight=False, size=None):
    r = par.add_run(s)
    r.italic = italic or None
    r.bold = bold or None
    if font:
        r.font.name = font
    if size:
        r.font.size = Pt(size)
    if sup:
        r.font.superscript = True
    if sub:
        r.font.subscript = True
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


# ------------------------------------------------------------------ helpers
def set_line_numbering(section):
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sectPr.append(ln)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)


def para(doc, text, style=None, align=None, space_after=6, italic=False, bold=False, size=None):
    p = doc.add_paragraph(style=style)
    add_inline(p, text, base_italic=italic, base_bold=bold, size=size)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def heading(doc, text, level):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12 if level == 1 else 12)
    r.italic = (level == 2)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, table):
    para(doc, table["caption"], bold=False, space_after=4)
    rows = table["rows"]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            add_inline(c.paragraphs[0], cell, base_bold=(i == 0), size=10)
    doc.add_paragraph()


def add_figure(doc, fig, width_cm=16.0):
    path = os.path.join(HERE, fig["path"])
    if not os.path.exists(path):
        raise SystemExit(f"missing figure {path}")
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, fig["caption"], space_after=12, size=10)


# ------------------------------------------------------------------- build
def build(md, out_docx, figures="end", line_numbers=True, spacing=1.5):
    d = parse(md)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    st.paragraph_format.line_spacing = spacing
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    if line_numbers:
        set_line_numbering(sec)
    add_page_number(sec)

    # title page block
    p = doc.add_paragraph(); r = p.add_run(d["title"]); r.bold = True; r.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(12)
    para(doc, d["meta"].get("AUTHOR", ""), space_after=2)
    para(doc, d["meta"].get("AFFIL", ""), space_after=2, italic=True)
    para(doc, "Corresponding author: " + d["meta"].get("EMAIL", ""), space_after=2)
    if d["meta"].get("ORCID"):
        para(doc, "ORCID: " + d["meta"]["ORCID"], space_after=2)
    doc.add_paragraph()

    fig_by_num = {int(re.match(r"Fig\. (\d+)\.", f["caption"]).group(1)): f for f in d["figures"]}
    placed = set()
    for kind, val in d["blocks"]:
        if kind == "h1":
            heading(doc, val, 1)
        elif kind == "h2":
            heading(doc, val, 2)
        elif kind == "keywords":
            p = para(doc, "", space_after=12)
            _run(p, "Keywords: ", bold=True); add_inline(p, val)
        elif kind == "table":
            add_table(doc, d["table"])
        elif kind == "p":
            para(doc, val)
            if figures == "inline":
                for n in sorted(fig_by_num):
                    if n not in placed and re.search(rf"Fig\. {n}\b", val):
                        add_figure(doc, fig_by_num[n]); placed.add(n)
    if figures == "end":
        heading(doc, "Figures", 1)
        for n in sorted(fig_by_num):
            add_figure(doc, fig_by_num[n])
    else:
        for n in sorted(fig_by_num):
            if n not in placed:
                add_figure(doc, fig_by_num[n])
    doc.core_properties.title = d["title"]
    doc.core_properties.author = d["meta"].get("AUTHOR", "")
    doc.save(out_docx)
    return d


def to_pdf(docx_path):
    outdir = os.path.dirname(docx_path)
    cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=300)
    return os.path.splitext(docx_path)[0] + ".pdf"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--figures", choices=["inline", "end"], default="end")
    ap.add_argument("--stem", default="manuscript")
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--no-line-numbers", action="store_true")
    ap.add_argument("--spacing", type=float, default=1.5)
    a = ap.parse_args()
    md = open(SRC, encoding="utf-8").read()
    out = os.path.join(HERE, a.stem + ".docx")
    d = build(md, out, figures=a.figures, line_numbers=not a.no_line_numbers, spacing=a.spacing)
    print("wrote", out)
    if a.pdf:
        print("wrote", to_pdf(out))
    words = len(re.findall(r"\S+", " ".join(v for k, v in d["blocks"] if k == "p")))
    abstract = next((v for k, v in d["blocks"] if k == "p"), "")
    print(f"body words (paragraphs incl. abstract/declarations): {words}; abstract words: {len(abstract.split())}")


if __name__ == "__main__":
    main()
